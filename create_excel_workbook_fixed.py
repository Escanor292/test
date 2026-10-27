import pandas as pd
import os
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows

# Đường dẫn
base_path = r"d:\Kiem_Thu\testing"
dest_path = r"d:\Kiem_Thu\testing\23CT113_Nhom02_NenTangCrowdfunding(TuTeFund)\Test Case_23CT113_Nhom02_COMPLETE.xlsx"

# Đọc tất cả CSV files với error handling
try:
    test_cases_df = pd.read_csv(os.path.join(base_path, "Test_Cases_Sheet.csv"), on_bad_lines='skip')
    print("Đã đọc Test_Cases_Sheet.csv")
except Exception as e:
    print(f"Lỗi đọc Test_Cases_Sheet.csv: {e}")
    test_cases_df = pd.DataFrame()

try:
    requirements_df = pd.read_csv(os.path.join(base_path, "Requirements_Sheet.csv"))
    print("Đã đọc Requirements_Sheet.csv")
except Exception as e:
    print(f"Lỗi đọc Requirements_Sheet.csv: {e}")
    requirements_df = pd.DataFrame()

try:
    traceability_df = pd.read_csv(os.path.join(base_path, "Traceability_Matrix_Sheet.csv"))
    print("Đã đọc Traceability_Matrix_Sheet.csv")
except Exception as e:
    print(f"Lỗi đọc Traceability_Matrix_Sheet.csv: {e}")
    traceability_df = pd.DataFrame()

try:
    summary_df = pd.read_csv(os.path.join(base_path, "Summary_Sheet.csv"))
    print("Đã đọc Summary_Sheet.csv")
except Exception as e:
    print(f"Lỗi đọc Summary_Sheet.csv: {e}")
    summary_df = pd.DataFrame()

try:
    readme_df = pd.read_csv(os.path.join(base_path, "Read_Me_Sheet.csv"))
    print("Đã đọc Read_Me_Sheet.csv")
except Exception as e:
    print(f"Lỗi đọc Read_Me_Sheet.csv: {e}")
    readme_df = pd.DataFrame()

# Tạo workbook với openpyxl để có nhiều sheets
wb = Workbook()

# Sheet 1: Read Me
ws = wb.active
ws.title = "Read Me"
if not readme_df.empty:
    for r_idx, row in enumerate(dataframe_to_rows(readme_df, index=False, header=True), 1):
        for c_idx, value in enumerate(row, 1):
            ws.cell(row=r_idx, column=c_idx, value=value)
else:
    ws['A1'] = "Read Me Sheet"
    ws['A2'] = "Workbook Test Case Documentation - Tử Tế Fund (23CT113_Nhom02)"

# Sheet 2: Requirements
ws2 = wb.create_sheet("Requirements")
if not requirements_df.empty:
    for r_idx, row in enumerate(dataframe_to_rows(requirements_df, index=False, header=True), 1):
        for c_idx, value in enumerate(row, 1):
            ws2.cell(row=r_idx, column=c_idx, value=value)
else:
    headers = ["Requirement ID", "Module", "Requirement Title", "Requirement Description", 
               "Stakeholder Role", "Acceptance Criteria", "Priority", "Source/Reference Requirement", 
               "Confirmation Status", "Requirement Type"]
    for c_idx, header in enumerate(headers, 1):
        ws2.cell(row=1, column=c_idx, value=header)

# Sheet 3: Test Cases
ws3 = wb.create_sheet("Test Cases")
if not test_cases_df.empty:
    for r_idx, row in enumerate(dataframe_to_rows(test_cases_df, index=False, header=True), 1):
        for c_idx, value in enumerate(row, 1):
            ws3.cell(row=r_idx, column=c_idx, value=value)
else:
    headers = ["Test Case ID", "Test Requirement ID", "Module", "Test Case Title", "Objective", 
               "Test Approach", "Test Type", "Test Technique", "Priority", "Precondition", 
               "Test Data", "Test Steps", "Expected Result", "Test Environment", "Execution Method", "Assignee"]
    for c_idx, header in enumerate(headers, 1):
        ws3.cell(row=1, column=c_idx, value=header)

# Sheet 4: Traceability Matrix
ws4 = wb.create_sheet("Traceability Matrix")
if not traceability_df.empty:
    for r_idx, row in enumerate(dataframe_to_rows(traceability_df, index=False, header=True), 1):
        for c_idx, value in enumerate(row, 1):
            ws4.cell(row=r_idx, column=c_idx, value=value)
else:
    headers = ["REQ ID", "TR ID", "Module", "Requirement Description", "Test Requirement Description", 
               "TC ID", "Test Case Description", "Test Coverage Status", "Coverage Percentage", "Test Status"]
    for c_idx, header in enumerate(headers, 1):
        ws4.cell(row=1, column=c_idx, value=header)

# Sheet 5: Summary
ws5 = wb.create_sheet("Summary")
if not summary_df.empty:
    for r_idx, row in enumerate(dataframe_to_rows(summary_df, index=False, header=True), 1):
        for c_idx, value in enumerate(row, 1):
            ws5.cell(row=r_idx, column=c_idx, value=value)
else:
    headers = ["Category", "Count", "Percentage", "Notes"]
    for c_idx, header in enumerate(headers, 1):
        ws5.cell(row=1, column=c_idx, value=header)

# Sheet 6: Test Execution (Empty template)
ws6 = wb.create_sheet("Test Execution")
execution_headers = ["Execution ID", "TC ID", "Lần chạy", "Ngày/giờ", "Người kiểm thử", 
                     "Môi trường thực tế", "Build/phiên bản", "Hệ điều hành", 
                     "Trình duyệt/phiên bản", "Công cụ", "Manual/Automated", 
                     "Actual Result", "Status", "Evidence", "Bug ID", "Ghi chú"]
for c_idx, header in enumerate(execution_headers, 1):
    ws6.cell(row=1, column=c_idx, value=header)

# Sheet 7: Bug Log (Empty template)
ws7 = wb.create_sheet("Bug Log")
bug_headers = ["Bug ID", "TC ID", "TR ID", "Module", "Tiêu đề", "Mô tả lỗi", 
               "Môi trường", "Build/phiên bản", "Mức độ", "Precondition", 
               "Steps to Reproduce", "Expected Result", "Actual Result", 
               "Evidence", "Trạng thái xử lý", "Kết quả Retest"]
for c_idx, header in enumerate(bug_headers, 1):
    ws7.cell(row=1, column=c_idx, value=header)

# Sheet 8: TR List (Test Requirements)
ws8 = wb.create_sheet("TR List")
tr_headers = ["TR ID", "REQ ID", "Module", "Mô tả yêu cầu kiểm thử", "Loại yêu cầu", 
              "Ưu tiên", "Tiêu chí chấp nhận", "Nguồn", "Trạng thái xác nhận"]
for c_idx, header in enumerate(tr_headers, 1):
    ws8.cell(row=1, column=c_idx, value=header)

# Lưu workbook
wb.save(dest_path)
print(f"\nWorkbook đã được tạo tại: {dest_path}")
print(f"Số sheets: {len(wb.sheetnames)}")
print(f"Tên sheets: {wb.sheetnames}")