# -*- coding: utf-8 -*-
"""Sinh Test Plan 23CT113_Nhom02 v3 tu ban FIXED.

Moi thay doi deu bam ma nguon https://github.com/Escanor292/platform (main, commit 2ce4b79)
va hanh vi da kiem chung tren ban trien khai ngay 29/10/2026.
"""
import copy
import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

SRC = "23CT113_Nhom02_NenTangCrowdfunding(TuTeFund)/Test Plan_23CT113_Nhom02_FIXED.docx"
OUT = "23CT113_Nhom02_NenTangCrowdfunding(TuTeFund)/Test Plan_23CT113_Nhom02_v3.docx"

d = docx.Document(SRC)


# ----------------------------------------------------------------------------- helpers
def set_par_text(par, text):
    """Giu dinh dang cua run dau tien, thay toan bo text."""
    if par.runs:
        par.runs[0].text = text
        for r in par.runs[1:]:
            r._element.getparent().remove(r._element)
    else:
        par.add_run(text)


def par_by_prefix(prefix, start=0):
    for i, p in enumerate(d.paragraphs):
        if i < start:
            continue
        if p.text.strip().startswith(prefix):
            return p
    raise KeyError(f"Khong tim thay doan bat dau bang: {prefix!r}")


def replace_par(prefix, text):
    p = par_by_prefix(prefix)
    set_par_text(p, text)
    return p


def delete_par(prefix):
    p = par_by_prefix(prefix)
    p._p.getparent().remove(p._p)


def insert_after(ref_par, texts):
    """Chen cac doan moi ngay sau ref_par, giu style cua ref_par."""
    anchor = ref_par
    for text in texts:
        new_p = copy.deepcopy(anchor._p)
        anchor._p.addnext(new_p)
        par = Paragraph(new_p, anchor._parent)
        set_par_text(par, text)
        anchor = par
    return anchor


def set_cell_text(cell, text):
    first = cell.paragraphs[0]
    set_par_text(first, text)
    for extra in cell.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)


def reshape_row(tr, ncols):
    tcs = tr.findall(qn("w:tc"))
    while len(tcs) < ncols:
        tr.append(copy.deepcopy(tcs[-1]))
        tcs = tr.findall(qn("w:tc"))
    while len(tcs) > ncols:
        tr.remove(tcs[-1])
        tcs = tr.findall(qn("w:tc"))
    for tc in tcs:
        tcPr = tc.find(qn("w:tcPr"))
        if tcPr is None:
            continue
        for tag in ("w:gridSpan", "w:vMerge"):
            el = tcPr.find(qn(tag))
            if el is not None:
                tcPr.remove(el)


def set_grid(tbl, ratios):
    """Dat lai tblGrid theo ti le cho truoc."""
    grid = tbl._tbl.find(qn("w:tblGrid"))
    cols = grid.findall(qn("w:gridCol"))
    total = sum(int(c.get(qn("w:w"))) for c in cols)
    for c in cols:
        grid.remove(c)
    s = sum(ratios)
    for r in ratios:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(int(total * r / s)))
        grid.append(gc)


def mark_tr(tr, header=False):
    trPr = tr.find(qn("w:trPr"))
    if trPr is None:
        trPr = OxmlElement("w:trPr")
        tr.insert(0, trPr)
    if trPr.find(qn("w:cantSplit")) is None:
        trPr.append(OxmlElement("w:cantSplit"))
    if header and trPr.find(qn("w:tblHeader")) is None:
        el = OxmlElement("w:tblHeader")
        el.set(qn("w:val"), "true")
        trPr.append(el)


def rebuild_table(idx, data, ratios):
    """Ghi lai toan bo bang idx; data[0] la dong tieu de."""
    tbl = d.tables[idx]
    ncols = len(data[0])
    trs = tbl._tbl.findall(qn("w:tr"))
    tpl_header = copy.deepcopy(trs[0])
    tpl_body = copy.deepcopy(trs[1]) if len(trs) > 1 else copy.deepcopy(trs[0])
    for tr in trs:
        tbl._tbl.remove(tr)

    new_trs = []
    for row in data:
        tpl = tpl_header if row is data[0] else tpl_body
        tr = copy.deepcopy(tpl)
        reshape_row(tr, ncols)
        tbl._tbl.append(tr)
        new_trs.append(tr)

    set_grid(tbl, ratios)
    for ri, row in enumerate(data):
        for ci, val in enumerate(row):
            set_cell_text(tbl.rows[ri].cells[ci], val)
        mark_tr(new_trs[ri], header=(ri == 0))


def polish_tables():
    """Lap lai dong tieu de khi bang sang trang, khong tach dong giua hai trang."""
    for tbl in d.tables:
        for ri, tr in enumerate(tbl._tbl.findall(qn("w:tr"))):
            mark_tr(tr, header=(ri == 0))


# ----------------------------------------------------------------------------- front matter
replace_par("Phiên bản: 2.0.0",
            "Phiên bản: 2.1.0 — cập nhật ngày 29/10/2026 sau khi đối chiếu lại mã nguồn thanh toán")
replace_par("Ngày lập: 28/10/2026",
            "Ngày lập: 28/10/2026 — ngày cập nhật: 29/10/2026")
replace_par("Nguồn dữ liệu chính:",
            "Nguồn dữ liệu chính: PRD 2.0.0, workbook Test Case 23CT113_Nhom02, Traceability Matrix, "
            "và mã nguồn github.com/Escanor292/platform (nhánh main, commit 2ce4b79 — 189 API route, "
            "93 trang, 63 model Prisma/PostgreSQL).")
replace_par("Ghi chú: các phần ngoài phạm vi Phase 1",
            "Ghi chú về phạm vi: bộ 75 Test Case hiện hành chỉ phủ 5 phân hệ Phase 1 (Login/Auth, Campaign, "
            "Reward, Blog, KYC). Mọi chức năng còn lại đã có chủ sở hữu và được xếp vào Phase 2, không tính "
            "vào 75 Test Case. Cơ chế thanh toán của bản demo là mô hình ngân hàng trung gian (BANK_ESCROW) "
            "kèm chế độ giả lập dành cho admin, đã được kiểm chứng trực tiếp trên mã nguồn và bản triển khai.")

# ----------------------------------------------------------------------------- 1. Mục tiêu
replace_par("- Kiểm tra 29 Requirement",
            "- Kiểm tra 29 Requirement, 33 Test Requirement và 75 Test Case của 5 phân hệ Phase 1 đã được "
            "xác minh trong workbook.")
replace_par("- Ghi rõ các phần tồn tại trong mã nguồn",
            "- Phân định rõ ranh giới Phase 1 và Phase 2: 13 hạng mục đã có trong mã nguồn nhưng chưa được cấp "
            "Test Case (Checkout/thanh toán, kho đồ, hoàn tiền, chế độ giả lập, Project Hub, báo cáo vi phạm, "
            "badge, chat, membership, Profile Studio, thuế, bằng cấp, quản trị hệ thống) — mỗi hạng mục có đúng "
            "một chủ sở hữu và được ghi ở mục 2.2.")

# ----------------------------------------------------------------------------- 3. Mô hình người dùng
replace_par("Creator: tạo campaign ở trạng thái DRAFT",
            "Creator chờ duyệt (CREATOR_PENDING): chỉ được tạo campaign ở trạng thái DRAFT, gửi yêu cầu duyệt "
            "và chỉnh sửa nội dung; chưa có quyền nhận ủng hộ.")
replace_par("Creator: tạo và quản lý chiến dịch",
            "Creator đã duyệt (CREATOR): quản lý chiến dịch, cập nhật reward, theo dõi tiến độ gây quỹ, lưu sao "
            "kê; chỉ campaign ở trạng thái ACTIVE và còn trong hạn mới nhận pledge.")
replace_par("Lưu ý về thanh toán:",
            "Lưu ý về thanh toán: bản demo không dùng cổng thanh toán trực tuyến. Mọi route PayOS, VNPay, MoMo, "
            "SePay và webhook thanh toán đều trả HTTP 410 với mã ESCROW_BANK_ONLY. Luồng hiện hành là chuyển khoản vào "
            "tài khoản ngân hàng trung gian (paymentProvider = BANK_ESCROW) rồi admin đối soát bằng tay; ngoài ra "
            "admin có thể bật chế độ giả lập để ghi nhận thanh toán với mã giao dịch tiền tố SIM-. Chi tiết ở mục "
            "5.4 và 6.4.")

# ----------------------------------------------------------------------------- 2.2 table
rebuild_table(1, [
    ["Phân hệ", "Đường dẫn / chức năng trong mã nguồn", "Trạng thái và chủ sở hữu"],
    ["Checkout và thanh toán (BANK_ESCROW, admin đối soát)",
     "/api/checkout-sessions, /api/payments/cart, /api/admin/pledges/[id]/settle, /api/thanh-toan/chuyen-khoan/[pledgeId]",
     "Phase 2 — cơ chế đã xác minh, chưa cấp TC. Chủ: Nguyễn Quách Phú Tài"],
    ["Chế độ giả lập và vận chuyển giả lập",
     "/api/admin/simulation, /api/orders/[pledgeId]/simulate, src/lib/payment/simulate-checkout.ts, simulate-shipping.ts",
     "Phase 2 — chưa cấp TC. Chủ: Nguyễn Quách Phú Tài"],
    ["Hoàn tiền và tài khoản trung gian",
     "/api/payments/refund, src/lib/payment/refund.ts, escrow.ts, escrow-account.ts",
     "Phase 2 — chưa cấp TC. Chủ: Nguyễn Quách Phú Tài"],
    ["Kho đồ và purchases", "/purchases, /api/purchases/digital",
     "Phase 2 — chưa cấp TC. Chủ: Nguyễn Quách Phú Tài"],
    ["Project Hub", "/api/projects, /dashboard/creator/projects",
     "Phase 2 — chưa cấp TC. Chủ: Trần Xuân Ân"],
    ["Báo cáo vi phạm", "/api/reports, /dashboard/admin/reports",
     "Phase 2 — chưa cấp TC. Chủ chính: Bùi Đặng Quốc Khánh; Trần Xuân Ân phối hợp ở luồng duyệt"],
    ["Badge", "/api/admin/badges, /dashboard/admin/badges",
     "Phase 2 — chưa cấp TC. Chủ: Bùi Đặng Quốc Khánh"],
    ["Chat và gọi", "/chat, /messages",
     "Phase 2 — chưa cấp TC. Chủ: Nguyễn Trung Tiến"],
    ["Membership trả phí định kỳ", "/api/memberships, /dashboard/creator/hoi-vien",
     "Phase 2 — chưa cấp TC. Chủ: Nguyễn Trung Tiến"],
    ["Profile Studio", "/api/profile/customization, /profile/[userId]/customize",
     "Phase 2 — chưa cấp TC. Chủ: Nguyễn Khánh Du"],
    ["Thuế và chứng từ", "/api/tax, /hoa-don/[number], /chung-tu/[code]",
     "Phase 2 — chưa cấp TC. Chủ: Nguyễn Khánh Du"],
    ["Bằng cấp và chứng chỉ", "/api/credentials, /dashboard/bang-cap",
     "Phase 2 — chưa cấp TC. Chủ: Nguyễn Khánh Du"],
    ["Quản trị hệ thống", "/api/admin/*",
     "Phase 2, trừ duyệt KYC và duyệt campaign thuộc Phase 1. Chủ: Nguyễn Quách Phú Tài"],
], [0.20, 0.44, 0.36])

replace_par("Ghi chú: các phân hệ ngoài 5 phân hệ cốt lõi",
            "Ghi chú: con số 75 Test Case chỉ áp dụng cho 5 phân hệ Phase 1. Các hạng mục ở bảng trên đã có chủ "
            "sở hữu và nằm trong backlog Phase 2 — không hạng mục nào được tính là “đã kiểm thử” chỉ vì xuất hiện "
            "trong mã nguồn. Khi mở Phase 2, mỗi hạng mục phải được cấp REQ/TR/TC riêng và cộng vào tổng của "
            "workbook, không tái sử dụng 75 Test Case hiện hành.")

# ----------------------------------------------------------------------------- 4.1 payment data
replace_par("- Luồng thanh toán được kiểm thử ở mức giả lập:",
            "- Luồng thanh toán: pledge ONLINE tạo bản ghi PENDING với paymentProvider = BANK_ESCROW và nội dung "
            "chuyển khoản tới tài khoản trung gian; pledge COD chỉ dành cho sản phẩm vật lý có sẵn và không cộng "
            "vào số gây quỹ. Dữ liệu giả lập do admin tạo ra luôn mang mã giao dịch tiền tố SIM- hoặc SIM: để có "
            "thể dọn sạch bằng DELETE /api/admin/simulation trước mỗi vòng kiểm thử.")

# ----------------------------------------------------------------------------- 4.2 data table
rebuild_table(2, [
    ["Phân hệ", "Dữ liệu cần chuẩn bị"],
    ["Login / Auth", "Email hợp lệ, email đã tồn tại, mật khẩu dưới/đúng biên, token đặt lại mật khẩu, role BACKER / CREATOR_PENDING / CREATOR / ADMIN"],
    ["Campaign", "Campaign ở trạng thái DRAFT, PENDING_REVIEW, ACTIVE, SUCCESS, FAILED, CANCELED; dữ liệu có/không có reward, fundingModel khác nhau"],
    ["Reward và tồn kho", "Reward có giá trị âm/0/hợp lệ, stock null/0/5, loại PHYSICAL/DIGITAL, số lượng hạn mức"],
    ["Blog", "Bài ở DRAFT, PENDING_REVIEW, PUBLISHED; publishedAt trong quá khứ/tương lai; payload chứa thẻ script hay ảnh lớn"],
    ["KYC", "Hồ sơ Pending, Verified, Rejected; ảnh JPG, PDF, sai định dạng, file quá dung lượng, video MP4"],
    ["Checkout (Phase 2)", "Pledge PENDING với paymentProvider BANK_ESCROW, COD; nội dung chuyển khoản và tài khoản trung gian; mã SIM- do admin tạo; trạng thái vận chuyển SHIPPED / DELIVERED / DELIVERY_FAILED / RETURNED"],
], [0.22, 0.78])

# ----------------------------------------------------------------------------- 4.3 accounts
rebuild_table(3, [
    ["Vai trò", "Điều kiện", "Mục đích kiểm thử"],
    ["Guest", "Chưa đăng nhập", "Xem chiến dịch công khai, blog public, tra cứu thông tin công khai"],
    ["Backer", "role = BACKER", "Đăng ký, ủng hộ, bình luận, nâng cấp Creator"],
    ["Creator chờ duyệt", "role = CREATOR_PENDING", "Kiểm tra giới hạn quyền khi chưa được duyệt: chỉ tạo DRAFT và gửi duyệt"],
    ["Creator đã duyệt", "role = CREATOR", "Tạo và quản lý campaign, reward, cập nhật tiến độ"],
    ["Admin", "role = ADMIN hoặc isAdmin = true", "Duyệt campaign, duyệt KYC, đối soát thanh toán, các thao tác quản trị"],
    ["Admin bật chế độ giả lập", "role = ADMIN và platform_settings.simulation_mode = true", "Ghi nhận thanh toán giả lập, giả lập vận chuyển, dọn dữ liệu giả lập (Phase 2)"],
], [0.24, 0.34, 0.42])

# ----------------------------------------------------------------------------- 5.1 environment
rebuild_table(4, [
    ["Môi trường", "Địa chỉ / công cụ", "Mục đích sử dụng"],
    ["Local", "http://localhost:3000 (npm run dev)", "Phát triển và kiểm thử hồi quy nhanh; chạy Jest/Playwright tại máy"],
    ["Staging chính", "https://platform-seven-navy-44.vercel.app", "Bản main hiện hành: đối chiếu UI/UX, kiểm cơ chế thanh toán và endpoint giả lập. HTTP 200, kiểm ngày 29/10/2026"],
    ["Staging cũ", "https://test-gamma-azure-65.vercel.app", "Bản build cũ: cổng thanh toán trả 410 nhưng chưa có endpoint giả lập (404). HTTP 200, kiểm ngày 29/10/2026"],
    ["Cơ sở dữ liệu", "PostgreSQL qua Prisma (schema.prisma, 63 model)", "Kiểm trạng thái dữ liệu sau thao tác: pledges, campaigns, users, kyc_info"],
    ["Cache", "Redis (ioredis), cache đọc và fail-open", "Kiểm dữ liệu hiển thị có thể trễ; xoá cache trước khi đối chiếu số huy động"],
    ["CI", "GitHub Actions .github/workflows/platform-ci.yml", "Job quality (prisma generate + Jest + tsc) và job e2e (Postgres 16 + Playwright Chromium)"],
    ["Repository", "https://github.com/Escanor292/platform.git (main, commit 2ce4b79)", "Đối chiếu phiên bản mã nguồn khi viết và chạy test case"],
], [0.17, 0.40, 0.43])

# ----------------------------------------------------------------------------- 5.2 platform
rebuild_table(5, [
    ["Hạng mục", "Cấu hình"],
    ["Hệ điều hành", "Windows 11 (64-bit)"],
    ["Trình duyệt chính", "Google Chrome phiên bản mới nhất"],
    ["Trình duyệt đối chiếu", "Microsoft Edge và Firefox"],
    ["Độ phân giải desktop", "1366×768 và 1920×1080"],
    ["Độ phân giải mobile", "390×844"],
    ["Phương pháp kiểm tra responsive", "Mô phỏng bằng Chrome DevTools Device Mode (không dùng thiết bị thật); mọi lỗi vỡ giao diện phải có ảnh chụp kèm kích thước viewport"],
    ["Mạng", "Internet bình thường và kịch bản làm chậm (DevTools throttling) để kiểm phản hồi"],
], [0.28, 0.72])

# ----------------------------------------------------------------------------- 5.3 tools
replace_par("- Tự động hóa: Jest và Playwright",
            "- Tự động hóa: repo có Jest (npm test, test:ui, test:functional, test:integration), Playwright e2e "
            "trên Chromium và CI platform-ci.yml gồm 2 job: quality (prisma generate + Jest + tsc) và e2e "
            "(Postgres 16, seed, Chromium). Bộ 75 Test Case trong workbook vẫn được thiết kế để chạy thủ công; "
            "phần tự động chỉ dùng để chặn hồi quy trước mỗi vòng kiểm thử.")
replace_par("- Quản lý kiểm thử: workbook Excel",
            "- Quản lý kiểm thử: workbook Excel gồm Requirements, Test Cases, Traceability Matrix, Test Execution, "
            "Bug Log và Summary; kết quả chỉ được ghi sau khi thực thi thật.")

# ----------------------------------------------------------------------------- 5.4 payment
replace_par("5.4 Ghi chú về thanh toán",
            "5.4 Cơ chế thanh toán của bản demo (đã kiểm chứng)")
p54 = par_by_prefix("Các luồng checkout/thanh toán giả lập")
set_par_text(p54, "Bản demo không nối cổng thanh toán trực tuyến. Toàn bộ hành vi dưới đây được đối chiếu với "
                  "mã nguồn platform (main, commit 2ce4b79) và kiểm trực tiếp trên bản triển khai ngày 29/10/2026:")
insert_after(p54, [
    "- Cổng trực tuyến đã tắt: /api/payment/payos/create, /vnpay/create, /momo/create, /sepay/create, "
    "/api/payment/payos/mock-checkout và /api/payments/webhook đều trả HTTP 410 với body "
    "{\"code\":\"ESCROW_BANK_ONLY\"} — đã gọi thử cả 6 route ngày 29/10/2026 trên platform-seven-navy-44. "
    "Đây là hành vi đúng, không phải lỗi.",
    "- Luồng thanh toán hiện hành là BANK_ESCROW: pledge ONLINE được tạo ở trạng thái PENDING với "
    "paymentProvider = BANK_ESCROW, kèm transferContent và tài khoản ngân hàng trung gian (bankName, "
    "accountNumber, accountHolder) lấy từ GET /api/thanh-toan/chuyen-khoan/[pledgeId]; người ủng hộ chuyển "
    "khoản hoặc quét VietQR.",
    "- Đối soát: admin gọi POST /api/admin/pledges/[id]/settle để chuyển pledge sang SUCCESS (settlePledgeAsPaid, "
    "lý do “Admin đối soát tiền vào tài khoản ngân hàng trung gian”). Gọi khi không có session admin trả 403 "
    "(đã kiểm chứng).",
    "- COD: chỉ áp dụng cho sản phẩm vật lý còn hàng; trường hợp khác trả 400 “COD chỉ áp dụng cho sản phẩm vật lý "
    "có sẵn”. Đơn COD mang paymentProvider = COD và không cộng vào số gây quỹ của chiến dịch.",
    "- Webhook ngân hàng: POST /api/payment/bank-transfer yêu cầu header Authorization; thiếu hoặc sai trả 401 "
    "{\"success\":false,\"message\":\"Unauthorized\"} (đã kiểm chứng).",
    "- Chế độ giả lập (chỉ admin): bật qua khoá platform_settings.simulation_mode trong /api/admin/settings. Khi bật, "
    "simulationApplies() cho phép admin ghi nhận thanh toán giả lập: pledge nhận transactionId tiền tố SIM- (đơn lẻ) "
    "hoặc SIM: (cả giỏ), status SUCCESS, paymentProvider = SIMULATION.",
    "- Dọn dữ liệu giả lập: DELETE /api/admin/simulation gọi wipeSimulatedPledges() để hoàn tiền các pledge SIM- đang "
    "SUCCESS, chuyển membership liên quan sang EXPIRED và xoá donation_certificates; trả về số bản ghi đã gỡ.",
    "- Giả lập vận chuyển: POST /api/orders/[pledgeId]/simulate (chỉ admin) chuyển SHIPPED → DELIVERED, "
    "DELIVERY_FAILED hoặc RETURNED, sinh mã vận đơn theo hãng (GHTK/GHN/VTP/SPX/JNT); khi DELIVERED sẽ gọi "
    "onPledgeSuccess và createBackerInvoice.",
    "- Lưu ý môi trường: hai endpoint giả lập nêu trên có trên bản platform-seven-navy-44 nhưng chưa có trên "
    "test-gamma-azure-65 (trả 404). Phải chốt môi trường trước khi chạy nhóm test case giả lập.",
])

# ----------------------------------------------------------------------------- 6.1 typo
replace_par("- Daily standup:",
            "- Daily standup: thống nhất trạng thái test case, lỗi đang chặn, dữ liệu test và thời gian thực hiện "
            "tiếp theo.")

# ----------------------------------------------------------------------------- 6.2 test types
rebuild_table(6, [
    ["Loại kiểm thử", "Mục đích", "Phạm vi áp dụng"],
    ["Smoke Test", "Xác minh build cơ bản còn hoạt động trước khi kiểm thử sâu", "Đăng nhập, xem campaign, tạo campaign, gửi duyệt, reward, blog, KYC"],
    ["Functional Test", "Kiểm tra luồng nghiệp vụ đúng theo yêu cầu", "Toàn bộ 5 phân hệ Phase 1"],
    ["Validation Test", "Kiểm tra ràng buộc dữ liệu và thông báo lỗi", "Đăng ký, password, upload KYC, reward, campaign"],
    ["Security Test", "Kiểm tra phân quyền và bảo vệ dữ liệu", "Admin/Creator/Backer; các route admin và settle"],
    ["UI Test", "Kiểm tra hiển thị và phản hồi giao diện", "Campaign list, blog detail, dashboard, KYC upload"],
    ["Regression Test", "Xác nhận lỗi sửa không tạo lỗi mới", "Các case mức High sau khi fix; chạy Jest và Playwright"],
    ["Exploratory Test", "Tìm lỗi ngoài kịch bản chuẩn", "5 phiên 60 phút theo charter, mỗi thành viên một phiên"],
    ["Integration Test (thanh toán)", "Kiểm tra ghép nối escrow, đối soát và webhook", "Phase 2: /api/payment/*, /api/admin/pledges/[id]/settle, /api/payment/bank-transfer"],
    ["State Transition Test (vận chuyển)", "Kiểm tra chuỗi trạng thái giao hàng", "Phase 2: POST /api/orders/[pledgeId]/simulate"],
], [0.22, 0.36, 0.42])

# ----------------------------------------------------------------------------- 6.4 payment testing
# go 3 gach dau dong cu (da duoc thay bang 8 nhom moi ben duoi)
for _stale in (
    "- Kiểm tra các trường hợp success, failed, invalid input",
    "- Phân biệt rõ logic thanh toán giả lập (mock)",
    "- Nếu cần tích hợp thật qua PayOS",
):
    delete_par(_stale)

replace_par("6.4 Kiểm thử chuyên biệt cho thanh toán giả lập",
            "6.4 Kiểm thử thanh toán và chế độ giả lập (Phase 2 — chưa cấp Test Case)")
p64 = par_by_prefix("- Kiểm tra dữ liệu đầu vào số tiền, số lượng reward")
set_par_text(p64, "- Cổng trực tuyến: với từng route payos, vnpay, momo, sepay và /api/payments/webhook, gọi cả GET "
                  "và POST đều phải trả 410 kèm code ESCROW_BANK_ONLY; không route nào được tạo giao dịch.")
insert_after(p64, [
    "- BANK_ESCROW: pledge ONLINE hợp lệ tạo trạng thái PENDING, paymentProvider = BANK_ESCROW, có transferContent "
    "và đủ 3 trường tài khoản trung gian; mở /thanh-toan/chuyen-khoan?pledge=... hiển thị đúng số tiền và nội dung.",
    "- Đối soát: admin settle chuyển pledge sang SUCCESS và cộng đúng số huy động; user thường, Creator và khách "
    "đều nhận 403.",
    "- Chống trùng: settle lại một pledge đã SUCCESS trả ok kèm skipped = already-settled, không sinh giao dịch "
    "mới và raisedAmount của campaign không đổi; pledge ở trạng thái khác trả not-settlable.",
    "- COD: sản phẩm vật lý còn hàng đặt được; sản phẩm số hoặc hết hàng trả 400; đơn COD không làm tăng số gây quỹ.",
    "- Webhook ngân hàng: thiếu hoặc sai Authorization trả 401; đúng khoá thì applyBankCredit ghi nhận tiền vào.",
    "- Chế độ giả lập: khi simulation_mode = false, admin gọi giả lập nhận 403 “Hãy bật giả lập trong quản trị”; khi "
    "bật, pledge sinh mã SIM- và mang paymentProvider = SIMULATION; sau DELETE /api/admin/simulation, pledge SIM- "
    "được hoàn tiền, membership chuyển EXPIRED, chứng từ bị xoá.",
    "- Vận chuyển giả lập: kiểm đủ 4 nhánh SHIPPED, DELIVERED, DELIVERY_FAILED, RETURNED; DELIVERED sinh hoá đơn và "
    "chứng từ; đơn không phải PHYSICAL trả 400.",
    "Ghi chú: 8 nhóm trên là thiết kế kiểm thử đã sẵn sàng cho Phase 2, hiện chưa được cấp Test Case và không tính "
    "vào 75 Test Case của Phase 1.",
])

# ----------------------------------------------------------------------------- 7.1 entry criteria
replace_par("5. Đã có đủ tài khoản",
            "5. Đã có đủ 5 loại tài khoản: Guest, Backer, Creator chờ duyệt, Creator đã duyệt và Admin.")
p66 = par_by_prefix("6. Dữ liệu kiểm thử đã được seed")
insert_after(p66, [
    "7. Với các test case thanh toán (Phase 2), môi trường phải có endpoint giả lập và trạng thái "
    "simulation_mode phải được ghi rõ trong cột Test Environment của từng case.",
])

# ----------------------------------------------------------------------------- 8.2 exit criteria
replace_par("- Tỷ lệ Pass toàn bộ đạt ≥ 90%",
            "- Tỷ lệ Pass đạt ≥ 90% tính trên số test case đã thực thi; tối thiểu phải đạt 68/75 test case Pass.")
replace_par("- Toàn bộ lỗi High/Critical đã được sửa",
            "- Toàn bộ lỗi Blocker/Critical ở luồng must-have phải được đóng trước nghiệm thu. Lỗi High chỉ được "
            "phép tồn tại khi có đánh giá rủi ro, người phê duyệt (Test Lead) và căn cứ chấp nhận ghi rõ trong "
            "Test Report; không áp dụng ngoại lệ cho Blocker/Critical.")
replace_par("- Test Report có đủ 4 phần",
            "- Test Report có đủ 4 phần: tóm tắt kết quả, phân tích lỗi, đánh giá chất lượng và khuyến nghị; số "
            "liệu Pass/Fail/Bug phải khớp với sheet Test Execution và Bug Log.")

# ----------------------------------------------------------------------------- 8.3 source note
replace_par("Số liệu 29 REQ, 33 TR và 75 TC",
            "Số liệu 29 REQ, 33 TR và 75 TC lấy từ workbook “Test Case_23CT113_Nhom02.xlsx” (sheet Summary, đối "
            "chiếu ngày 29/10/2026: Requirements 29 dòng, Test Cases 75 dòng, Traceability Matrix 75 dòng). Các "
            "khẳng định về mã nguồn trong tài liệu này bám theo nhánh main của github.com/Escanor292/platform ở "
            "commit 2ce4b79; hành vi HTTP 410 của cổng thanh toán được kiểm trực tiếp trên bản triển khai.")

# ----------------------------------------------------------------------------- 9.1 schedule
rebuild_table(9, [
    ["Giai đoạn", "Thời gian", "Nội dung công việc", "Kết quả bàn giao", "Người phụ trách", "Trạng thái"],
    ["Khởi tạo", "07/09 – 13/09/2026", "Phân tích PRD, đối chiếu mã nguồn, chốt phạm vi 5 phân hệ", "PRD 2.0.0 và workbook ban đầu", "Nguyễn Quách Phú Tài", "Hoàn thành"],
    ["Thiết kế yêu cầu", "14/09 – 20/09/2026", "Lập 29 REQ và 33 TR từ PRD và mã nguồn", "Sheet Requirements và Traceability", "Cả nhóm, Test Lead duyệt", "Hoàn thành"],
    ["Thiết kế Test Case", "21/09 – 04/10/2026", "Viết 75 TC với Test Steps, Test Data, Expected Result", "75 TC trong workbook (15 TC/người)", "Cả nhóm", "Hoàn thành"],
    ["Chuẩn bị dữ liệu", "05/10 – 11/10/2026", "Seed dữ liệu, tạo tài khoản, dựng môi trường local và deploy", "Bộ dữ liệu theo mục 4", "Bùi Đặng Quốc Khánh", "Hoàn thành"],
    ["Rà soát Test Plan và đồng bộ mã nguồn", "28/10 – 29/10/2026", "Đối chiếu cơ chế thanh toán, sửa phạm vi, cập nhật lịch và tiêu chí", "Test Plan 2.1.0", "Nguyễn Quách Phú Tài", "Hoàn thành"],
    ["Thực thi vòng 1", "29/10 – 05/11/2026", "Chạy 75 TC, ghi Actual Result, Test Status và Bug Log", "Sheet Test Execution và Bug Log", "Cả nhóm", "Chưa thực thi (0/75)"],
    ["Sửa lỗi và retest", "06/11 – 09/11/2026", "Retest các case Fail, xác nhận fix, đóng lỗi", "Bug Log cập nhật trạng thái", "Cả nhóm", "Chưa thực thi"],
    ["Tổng kết", "10/11 – 11/11/2026", "Đối chiếu traceability, viết Test Report, chốt metrics", "Test Report và Metrics", "Nguyễn Khánh Du", "Chưa thực thi"],
], [0.17, 0.13, 0.29, 0.17, 0.15, 0.09])
set_grid(d.tables[9], [0.17, 0.13, 0.29, 0.17, 0.15, 0.09])

replace_par("2.2 Ngoài phạm vi kiểm thử hiện tại",
            "2.2 Phase 2 — hạng mục đã có trong mã nguồn nhưng ngoài bộ 75 Test Case hiện hành")

replace_par("9.3 Ghi chú về báo cáo vi phạm",
            "9.3 Chủ sở hữu báo cáo vi phạm")
replace_par("Báo cáo vi phạm là hạng mục phối hợp",
            "Báo cáo vi phạm thuộc Phase 2 và có đúng một chủ sở hữu chính: Bùi Đặng Quốc Khánh. Trần Xuân Ân "
            "phối hợp ở luồng admin duyệt campaign, vì thao tác duyệt báo cáo nằm trong màn quản trị campaign. "
            "Cách phân định: mọi test case kiểm nội dung và trạng thái báo cáo do Bùi Đặng Quốc Khánh viết; test "
            "case kiểm quyền duyệt campaign do Trần Xuân Ân viết và tham chiếu sang mã báo cáo, không viết trùng.")

# ----------------------------------------------------------------------------- 9.2 assignment
rebuild_table(10, [
    ["Thành viên", "Phase 1 (có Test Case)", "Số TC Phase 1", "Phase 2 và vai trò khác"],
    ["Nguyễn Quách Phú Tài", "Login / Auth / Hồ sơ", "15", "Phase 2: Checkout và thanh toán, kho đồ, hoàn tiền, chế độ giả lập. Test Lead, duyệt phạm vi, CI"],
    ["Trần Xuân Ân", "Campaign", "15", "Phase 2: Project Hub, phối hợp duyệt báo cáo vi phạm. Ma trận truy vết, tổng hợp Bug Log"],
    ["Nguyễn Khánh Du", "KYC", "15", "Phase 2: Profile Studio, bằng cấp, thuế và chứng từ. Test Report, tổng hợp Pass/Fail"],
    ["Bùi Đặng Quốc Khánh", "Reward và tồn kho", "15", "Phase 2: Badge, báo cáo vi phạm (chủ chính). Dữ liệu test, phân loại Severity"],
    ["Nguyễn Trung Tiến", "Blog", "15", "Phase 2: Chat và gọi, membership, UI public. Playwright, slide thuyết trình"],
    ["Tổng", "5 phân hệ", "75", "13 hạng mục Phase 2 đã có chủ, chưa cấp TC"],
], [0.19, 0.23, 0.11, 0.47])

# ----------------------------------------------------------------------------- 10. modules
replace_par("Authentication & profile:",
            "Authentication và profile (Phase 1): đăng ký, đăng nhập bằng email/mật khẩu hoặc Google, đổi mật khẩu, "
            "lưu hồ sơ, cập nhật avatar, kiểm soát quyền người dùng và nâng cấp Creator.")
replace_par("Campaign management:",
            "Campaign management (Phase 1): tạo campaign, cập nhật chi tiết, gửi duyệt, admin duyệt/từ chối, hủy "
            "chiến dịch và kiểm tra hoàn tiền theo mô hình quy định.")
replace_par("Reward & inventory:",
            "Reward và inventory (Phase 1): quản lý gói phần thưởng, kiểm tra điều kiện số lượng, giá trị, trạng thái "
            "tồn kho và logic hiển thị.")
replace_par("KYC & verification:",
            "KYC và verification (Phase 1): nộp hồ sơ, kiểm tra định dạng ảnh/tệp, trạng thái chờ/duyệt/từ chối và "
            "thông báo kết quả.")
replace_par("Blog & community:",
            "Blog và community (Phase 1): soạn bài, lưu nháp, duyệt, cập nhật trạng thái và quyền hiển thị bài viết.")
replace_par("Các module ngoài 5 phân hệ chính",
            "Checkout và thanh toán (Phase 2): cổng trực tuyến trả 410 ESCROW_BANK_ONLY; pledge BANK_ESCROW chờ admin "
            "đối soát; COD cho hàng vật lý có sẵn; webhook ngân hàng yêu cầu Authorization; chế độ giả lập SIM- do "
            "admin bật và dọn bằng DELETE /api/admin/simulation; vận chuyển giả lập 4 trạng thái. Các hạng mục Phase 2 "
            "còn lại (kho đồ, hoàn tiền, Project Hub, báo cáo vi phạm, badge, chat, membership, Profile Studio, thuế, "
            "bằng cấp, quản trị hệ thống) đã có chủ sở hữu theo mục 2.2 nhưng chưa được cấp Test Case.")

# ----------------------------------------------------------------------------- 11. risks
rebuild_table(11, [
    ["Mã", "Rủi ro", "Khả năng", "Ảnh hưởng", "Biện pháp giảm thiểu", "Người theo dõi"],
    ["R-01", "Môi trường online không ổn định hoặc phản hồi chậm", "Trung bình", "Cao", "Chuyển sang local khi online gián đoạn; ghi rõ môi trường trong Test Environment", "Nguyễn Quách Phú Tài"],
    ["R-02", "Dữ liệu kiểm thử bị thay đổi giữa các vòng", "Cao", "Cao", "Seed lại trước mỗi vòng, ghi lại ID bản ghi và trạng thái dùng test", "Bùi Đặng Quốc Khánh"],
    ["R-03", "KYC yêu cầu ảnh giấy tờ thật hoặc file vượt chuẩn", "Trung bình", "Cao", "Dùng dữ liệu giả lập, xác minh định dạng và kích thước tối đa", "Nguyễn Khánh Du"],
    ["R-04", "PRD hoặc mã nguồn thay đổi nhanh hơn Test Plan", "Cao", "Cao", "Đối chiếu PRD 2.0.0 và commit main mỗi vòng; cập nhật traceability cùng lúc", "Nguyễn Quách Phú Tài"],
    ["R-05", "Cổng thanh toán trả 410 bị ghi nhầm thành lỗi", "Cao", "Trung bình", "Ghi rõ ESCROW_BANK_ONLY là hành vi đúng trong Expected Result của case tương ứng", "Nguyễn Quách Phú Tài"],
    ["R-06", "Chế độ giả lập bật khi kiểm dữ liệu thật khiến pledge SIM- lẫn vào số huy động", "Trung bình", "Cao", "Chỉ bật khi chạy case giả lập; dọn bằng DELETE /api/admin/simulation trước và sau phiên; đối chiếu raisedAmount", "Nguyễn Quách Phú Tài"],
    ["R-07", "Redis cache còn dữ liệu cũ sau khi settle", "Trung bình", "Trung bình", "Xoá cache trước khi đối chiếu số huy động; ghi key cache vào Bug Log khi lệch số", "Bùi Đặng Quốc Khánh"],
    ["R-08", "Bản staging đang dùng thiếu endpoint giả lập", "Cao", "Trung bình", "Chốt môi trường platform-seven-navy-44 cho nhóm case giả lập; ghi rõ URL trong Test Environment", "Nguyễn Quách Phú Tài"],
], [0.05, 0.29, 0.09, 0.09, 0.33, 0.15])

replace_par("- Luồng thanh toán giả lập được coi là mock",
            "- Bản demo dùng mô hình BANK_ESCROW: không có giao dịch qua cổng trực tuyến, mọi route cổng trả 410; "
            "tiền được đối soát thủ công bởi admin hoặc ghi nhận qua chế độ giả lập với mã SIM-.")
replace_par("- Các phân hệ ngoài 5 phân hệ Phase 1",
            "- Các hạng mục Phase 2 đã có chủ sở hữu nhưng chưa có Test Case; chúng không được tính vào 75 Test Case "
            "và không dùng để kết luận nghiệm thu Phase 1.")
p115 = par_by_prefix("- Môi trường local và online dùng cùng")
insert_after(p115, [
    "- Redis chỉ đóng vai trò cache đọc và fail-open; dữ liệu hiển thị có thể trễ vài giây so với cơ sở dữ liệu.",
])

# ----------------------------------------------------------------------------- 12. conclusion
replace_par("Test Plan này xác định phạm vi kiểm thử cho 5 phân hệ",
            "Test Plan này xác định phạm vi kiểm thử Phase 1 cho 5 phân hệ — Login (Auth), Campaign, Reward, Blog và "
            "KYC — với 29 Requirement, 33 Test Requirement và 75 Test Case, mỗi thành viên 15 case. Cơ chế thanh "
            "toán của bản demo đã được làm rõ: không dùng cổng trực tuyến, pledge ONLINE đi theo BANK_ESCROW và admin "
            "đối soát, COD chỉ cho hàng vật lý có sẵn, còn chế độ giả lập SIM- và vận chuyển giả lập là công cụ của "
            "admin cho mục đích trình diễn. Nhóm case thanh toán được ghi nhận là Phase 2 với thiết kế đã sẵn sàng "
            "nhưng chưa cấp Test Case, nhờ vậy tài liệu không còn mâu thuẫn giữa phạm vi, phân công và số lượng test "
            "case. Mọi số liệu trong tài liệu bám theo workbook hiện hành và nhánh main của mã nguồn tại commit 2ce4b79.")

# ----------------------------------------------------------------------------- final
polish_tables()
d.save(OUT)
print("Đã ghi:", OUT)
