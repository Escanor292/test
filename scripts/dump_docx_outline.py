import sys, os
from docx import Document

FILES = {
    "OLD_NHOM2": r"23CT113_Nhom02_NenTangCrowdfunding(TuTeFund)/drive-download-20261027T100311Z-1-001/[Nhom2]_Ke_hoach_Kiem_thu.docx",
    "NEW_23CT113": r"23CT113_Nhom02_NenTangCrowdfunding(TuTeFund)/Test Plan_23CT113_Nhom02.docx",
    "REF_22CT115": r"22CT115_Nhom07_QuanLyBanHang(Sweet-Sip)/Test Plan_22CT115_Nhóm07.docx",
}

for key, path in FILES.items():
    if not os.path.exists(path):
        print(f"### {key}: MISSING {path}")
        continue
    d = Document(path)
    paras = d.paragraphs
    words = sum(len(p.text.split()) for p in paras)
    print(f"\n{'='*78}\n### {key}")
    print(f"paragraphs={len(paras)}  tables={len(d.tables)}  words(paras)={words}")
    # headings by style
    print("--- HEADINGS ---")
    for p in paras:
        st = p.style.name if p.style else ""
        if st.lower().startswith("heading") or st.lower().startswith("title"):
            t = p.text.strip()
            if t:
                print(f"  [{st}] {t}")
    print("--- TABLES ---")
    for i, tb in enumerate(d.tables):
        hdr = " | ".join(c.text.strip().replace("\n", " ")[:28] for c in tb.rows[0].cells)
        print(f"  T{i}: {len(tb.rows)}x{len(tb.columns)}  HEADER: {hdr}")
