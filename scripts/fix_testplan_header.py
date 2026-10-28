import copy, os, sys
from docx import Document

D = r"23CT113_Nhom02_NenTangCrowdfunding(TuTeFund)"
TARGET = os.path.join(D, "Test Plan_23CT113_Nhom02.docx")
LOCK = os.path.join(D, "~$st Plan_23CT113_Nhom02.docx")

if os.path.exists(LOCK):
    print("LOCKED: file dang mo trong Word ->", LOCK)
    sys.exit(2)

def set_text(p, text):
    """Giu dinh dang run dau tien, thay toan bo text."""
    runs = p.runs
    if not runs:
        r = p.add_run(text)
        return
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""

doc = Document(TARGET)
paras = doc.paragraphs

changed = []
def replace_prefix(prefix, new_text):
    for p in paras:
        if p.text.strip().startswith(prefix):
            set_text(p, new_text)
            changed.append(new_text)
            return p
    return None

replace_prefix("Dự án:", "Dự án: Nền tảng Gây quỹ Cộng đồng Tử Tế Fund (crowdfunding-vn)")
replace_prefix("Phiên bản:", "Phiên bản: 2.0.0")
replace_prefix("Ngày lập:", "Ngày lập: 07/09/2026")
anchor = replace_prefix("Người lập:", "Người lập: Nguyễn Quách Phú Tài (123000609) — Test Lead / Product Owner")

EXTRA = [
    "Mã tài liệu: TP-23CT113-Nhom02",
    "Nhóm: 23CT113 — Nhóm 02",
    "Căn cứ: PRD 2.0.0 (01-IC-Product-Requirements_23CT113_Nhom02.md) và mã nguồn hệ thống",
    "Phạm vi: 5 phân hệ — Login (Auth), Campaign, Reward, Blog, KYC",
]

if anchor is not None:
    cur = anchor._p
    for line in EXTRA:
        new_p = copy.deepcopy(anchor._p)
        cur.addnext(new_p)
        cur = new_p
        from docx.text.paragraph import Paragraph
        set_text(Paragraph(new_p, anchor._parent), line)
        changed.append(line)

doc.save(TARGET)
print("SAVED:", TARGET)
for c in changed:
    print("  -", c)
