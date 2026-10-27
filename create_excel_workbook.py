import pandas as pd
import os
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows

# Đường dẫn
base_path = r"d:\Kiem_Thu\testing"
dest_path = r"d:\Kiem_Thu\testing\23CT113_Nhom02_NenTangCrowdfunding(TuTeFund)\Test Case_23CT113_Nhom02_COMPLETE.xlsx"

# Đọc tất cả CSV files
test_cases_df = pd.read_csv(os.path.join(base_path, "Test_Cases_Sheet.csv"))
requirements_df = pd.read_csv(os.path.join(base_path, "Requirements_Sheet.csv"))
traceability_df = pd.read_csv(os.path.join(base_path, "Traceability_Matrix_Sheet.csv"))
summary_df = pd.read_csv(os.path.join(base_path, "Summary_Sheet.csv"))
readme_df = pd.read_csv(os.path.join(base_path, "Read_Me_Sheet.csv"))

# Tạo workbook với openpyxl để có nhiều sheets
wb = Workbook()

# Sheet 1: Read Me
ws = wb.active
ws.title = "Read Me"
for r_idx, row in enumerate(dataframe_to_rows(readme_df, index=False, header=True), 1):
    for c_idx, value in enumerate(row, 1):
        ws.cell(row=r_idx, column=c_idx, value=value)

# Sheet 2: Requirements
ws2 = wb.create_sheet("Requirements")
for r_idx, row in enumerate(dataframe_to_rows(requirements_df, index=False, header=True), 1):
    for c_idx, value in enumerate(row, 1):
        ws2.cell(row=r_idx, column=c_idx, value=value)

# Sheet 3: Test Cases
ws3 = wb.create_sheet("Test Cases")
for r_idx, row in enumerate(dataframe_to_rows(test_cases_df, index=False, header=True), 1):
    for c_idx, value in enumerate(row, 1):
        ws3.cell(row=r_idx, column=c_idx, value=value)

# Sheet 4: Traceability Matrix
ws4 = wb.create_sheet("Traceability Matrix")
for r_idx, row in enumerate(dataframe_to_rows(traceability_df, index=False, header=True), 1):
    for c_idx, value in enumerate(row, 1):
        ws4.cell(row=r_idx, column=c_idx, value=value)

# Sheet 5: Summary
ws5 = wb.create_sheet("Summary")
for r_idx, row in enumerate(dataframe_to_rows(summary_df, index=False, header=True), 1):
    for c_idx, value in enumerate(row, 1):
        ws5.cell(row=r_idx, column=c_idx, value=value)

# Sheet 6: Test Execution (Empty template)
ws6 = wb.create_sheet("Test Execution")
execution_headers = ["Execution ID", "TC ID", "Lần chạy", "Ngày/giờ", "Người kiểm thử", 
                     "Môi trường thực tế", "Build/phiên bản", "Hệ điều hành", 
                     "Trình duyệt/phiên bản", "Công cụ", "Manual/Automated", 
                     "Actual Result", "Status", "Evidence", "Bug ID", "Ghi chú"]
for c_idx, header in enumerate(execution_headers, 1):
    ws6.cell(row=1, column=c_idx, value=header)

# Sheet 7: Bug Log (Empty template)
ws7 = wb.create_sheet("Bug Log")
bug_headers = ["Bug ID", "TC ID", "TR ID", "Module", "Tiêu đề", "Mô tả lỗi", 
               "Môi trường", "Build/phiên bản", "Mức độ", "Precondition", 
               "Steps to Reproduce", "Expected Result", "Actual Result", 
               "Evidence", "Trạng thái xử lý", "Kết quả Retest"]
for c_idx, header in enumerate(bug_headers, 1):
    ws7.cell(row=1, column=c_idx, value=header)

# Lưu workbook
wb.save(dest_path)
print(f"Workbook đã được tạo tại: {dest_path}")